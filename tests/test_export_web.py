"""Excel export and the FastAPI endpoints."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from radar import web
from radar.export import IT_HEADERS, RADAR_HEADERS, export_workbook
from radar.models import Classification, Organization, Procedure
from radar.renewal import compute_renewals

NOW = datetime(2026, 9, 13, tzinfo=UTC)


def seed_radar(session, source: str = "synthetic") -> None:
    org = Organization(stir="301234567", name_canonical="Test Org (SYNTHETIC)",
                       region="Toshkent shahri")
    session.add(org)
    session.flush()
    proc = Procedure(source=source, source_id="X1", title="Fortinet FortiGate litsenziya 1 yil",
                     customer_org_id=org.id, completed_at=datetime(2025, 11, 17, tzinfo=UTC),
                     start_price=Decimal("48000000"), currency="UZS",
                     source_url="https://example.invalid/X1")
    session.add(proc)
    session.flush()
    session.add(Classification(procedure_id=proc.id, is_it=True, category="Firewall",
                               brand="Fortinet", is_subscription=True, term_months=12,
                               method="rule", confidence=0.9))
    session.commit()
    compute_renewals(session, now=NOW)


def test_workbook_structure_and_rows(session, tmp_path):
    seed_radar(session)
    path, count = export_workbook(session, tmp_path / "radar.xlsx", now=NOW)
    assert path.is_file() and count == 1
    book = load_workbook(path)
    assert book.sheetnames == ["Radar", "IT_Lots", "Export_10col", "Stats"]

    radar = book["Radar"]
    # Row 1 warns about synthetic sources, row 3 is the header, row 4 is data.
    assert "WARNING" in str(radar["A1"].value)
    assert [c.value for c in radar[3]] == RADAR_HEADERS
    row = [c.value for c in radar[4]]
    assert row[0] == "Test Org (SYNTHETIC)" and row[1] == "301234567"
    assert row[3] == "Firewall" and row[4] == "Fortinet"
    assert row[6] == 48000000.0 and isinstance(row[9], int)
    assert row[10] == "https://example.invalid/X1"

    lots = book["IT_Lots"]
    assert [c.value for c in lots[1]] == IT_HEADERS
    assert [c.value for c in lots[2]][:2] == ["synthetic", "X1"]

    stats = {r[0].value: r[1].value for r in book["Stats"].iter_rows(min_row=2)}
    assert stats["radar_rows"] == 1 and stats["it_lots"] == 1
    assert "synthetic" in str(stats["data_note"])


def test_workbook_without_synthetic_sources_has_no_warning(session, tmp_path):
    seed_radar(session, source="uzex")
    path, _ = export_workbook(session, tmp_path / "radar.xlsx", now=NOW)
    radar = load_workbook(path)["Radar"]
    assert [c.value for c in radar[1]] == RADAR_HEADERS


def test_empty_export_still_writes_a_file(session, tmp_path):
    path, count = export_workbook(session, tmp_path / "empty.xlsx", now=NOW)
    assert path.is_file() and count == 0


@pytest.fixture
def client(session, engine, monkeypatch):
    from contextlib import contextmanager

    @contextmanager
    def scope(*args, **kwargs):
        yield session

    monkeypatch.setattr(web, "session_scope", scope)
    monkeypatch.setattr(web, "get_engine", lambda *a, **k: engine)
    return TestClient(web.app)


def test_health_endpoint(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "database": "ok"}


def test_index_lists_endpoints(client):
    assert client.get("/").json() == {"endpoints": ["/radar", "/health"]}


def test_radar_page_renders_rows(client, session):
    seed_radar(session)
    body = client.get("/radar").text
    assert "Renewal Radar" in body
    assert "Test Org (SYNTHETIC)" in body
    assert "Fortinet" in body
    assert "301234567" in body
    assert "Synthetic data." in body, "synthetic rows must be labelled on the page"


def test_radar_page_empty_state(client):
    body = client.get("/radar").text
    assert "No opportunities in the contact window" in body
    assert "Synthetic data." not in body


def test_radar_page_hides_the_banner_for_real_sources(client, session):
    seed_radar(session, source="uzex")
    body = client.get("/radar").text
    assert "Test Org (SYNTHETIC)" in body
    assert "Synthetic data." not in body


def test_health_reports_a_broken_database(monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("no database")

    monkeypatch.setattr(web, "get_engine", boom)
    response = TestClient(web.app, raise_server_exceptions=False).get("/health")
    assert response.status_code == 503 and response.json()["status"] == "error"


def test_contact_now_counter(client, session):
    seed_radar(session)
    body = client.get("/radar").text
    assert "CONTACT NOW" in body.upper()


def test_radar_limit_is_applied(client, session):
    seed_radar(session)
    assert client.get("/radar?limit=0").status_code == 200


def test_lifecycle_window_is_reported(client):
    body = client.get("/radar").text
    assert "Contact window 60 days" in body
    assert "0&ndash;75" in body or "0–75" in body


def test_old_rows_drop_off(client, session):
    seed_radar(session)
    session.execute(Procedure.__table__.update().values(
        completed_at=NOW - timedelta(days=3000)))
    session.commit()
    compute_renewals(session, now=NOW)
    assert "Test Org (SYNTHETIC)" not in client.get("/radar").text


def test_api_search_and_eimzo_endpoints(client, session):
    seed_radar(session)

    # 1. Search keyword
    res = client.get("/api/search?q=Fortinet&type=keyword")
    assert res.status_code == 200
    data = res.json()
    assert "results" in data
    assert len(data["results"]) >= 1
    assert data["results"][0]["customer"]["name"] == "Test Org (SYNTHETIC)"

    # 2. Search customer
    res_cust = client.get("/api/search?q=301234567&type=customer")
    assert res_cust.status_code == 200
    cust_data = res_cust.json()
    assert cust_data["found"] is True
    assert cust_data["organizations"][0]["stir"] == "301234567"

    # 3. E-IMZO status endpoint
    res_eimzo = client.get("/api/eimzo/status")
    assert res_eimzo.status_code == 200
    assert "daemon" in res_eimzo.json()
    assert "session" in res_eimzo.json()

    # 4. Index endpoint lists endpoints
    res_index = client.get("/")
    assert res_index.status_code == 200
    assert "/radar" in res_index.json()["endpoints"]


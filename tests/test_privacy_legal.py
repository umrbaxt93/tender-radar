"""Tests for Phase 2: Privacy, PII Masking, AI Sanitization, and Data Retention."""

from __future__ import annotations

import io
import os
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
from starlette.testclient import TestClient

from radar.auth import SESSION_COOKIE_NAME, create_session_token, hash_password
from radar.cleanup import cleanup_old_snapshots, parse_retention_period
from radar.db import session_scope
from radar.models import (
    AuditLog,
    Classification,
    Organization,
    Procedure,
    RawSnapshot,
    RenewalOpportunity,
    User,
)
from radar.privacy import is_jshshir, is_stir, mask_identifier, sanitize_for_ai
from radar.web import app

client = TestClient(app)


# -----------------------------------------------------------------------------
# Fixtures
# -----------------------------------------------------------------------------
@pytest.fixture
def admin_user():
    username = f"admin_{os.urandom(4).hex()}"
    pwd = "AdminPass123!"
    with session_scope() as session:
        user = User(
            username=username,
            password_hash=hash_password(pwd),
            role="admin",
            is_active=True,
        )
        session.add(user)
        session.commit()
        user_id = user.id

    yield {"id": user_id, "username": username, "role": "admin"}

    with session_scope() as session:
        u = session.get(User, user_id)
        if u:
            session.delete(u)
            session.commit()


@pytest.fixture
def sales_user():
    username = f"sales_{os.urandom(4).hex()}"
    pwd = "SalesPass123!"
    with session_scope() as session:
        user = User(
            username=username,
            password_hash=hash_password(pwd),
            role="sales",
            is_active=True,
        )
        session.add(user)
        session.commit()
        user_id = user.id

    yield {"id": user_id, "username": username, "role": "sales"}

    with session_scope() as session:
        u = session.get(User, user_id)
        if u:
            session.delete(u)
            session.commit()


@pytest.fixture
def admin_cookie(admin_user):
    token = create_session_token(admin_user["id"], admin_user["username"], "admin")
    return {SESSION_COOKIE_NAME: token}


@pytest.fixture
def sales_cookie(sales_user):
    token = create_session_token(sales_user["id"], sales_user["username"], "sales")
    return {SESSION_COOKIE_NAME: token}


# -----------------------------------------------------------------------------
# 1. Identifier Detection & Masking Tests
# -----------------------------------------------------------------------------
def test_jshshir_vs_stir_detection():
    """Verify distinction between 9-digit corporate STIR and 14-digit individual JSHSHIR."""
    stir_legal = "308904387"
    jshshir_ind = "32405891230018"

    assert is_stir(stir_legal) is True
    assert is_jshshir(stir_legal) is False

    assert is_stir(jshshir_ind) is False
    assert is_jshshir(jshshir_ind) is True

    # Formatted with spaces / dashes
    assert is_stir(" 308-904-387 ") is True
    assert is_jshshir(" 324 058 912 300 18 ") is True

    # Invalid lengths
    assert is_stir("12345678") is False
    assert is_stir("1234567890") is False
    assert is_jshshir("1234567890123") is False
    assert is_jshshir("123456789012345") is False
    assert is_stir(None) is False
    assert is_jshshir("") is False


def test_mask_identifier_behavior():
    """Verify masking rules: 9-digit unmasked, 14-digit masked as *********12345 unless admin."""
    stir_legal = "308904387"
    jshshir_ind = "32405891230018"

    # 9-digit is always unmasked
    assert mask_identifier(stir_legal, is_admin=False) == "308904387"
    assert mask_identifier(stir_legal, is_admin=True) == "308904387"

    # 14-digit is masked for non-admin
    assert mask_identifier(jshshir_ind, is_admin=False) == "*********30018"

    # 14-digit is visible for admin
    assert mask_identifier(jshshir_ind, is_admin=True) == "32405891230018"

    # Empty / None handling
    assert mask_identifier(None) == ""
    assert mask_identifier("") == ""


# -----------------------------------------------------------------------------
# 2. AI Prompt Sanitization Tests
# -----------------------------------------------------------------------------
def test_sanitize_for_ai():
    """Verify PII (STIR, JSHSHIR, phone, email) is stripped before LLM ingestion."""
    text = (
        "Mijoz: YaTT Karimov A.B., JSHSHIR: 32405891230018. "
        "Kompaniya STIR: 308904387. Aloqa uchun tel: +998 (90) 123-45-67, "
        "email: info@supplier-uz.com. Mahsulot: 50 dona Windows Server litsenziyasi."
    )
    sanitized = sanitize_for_ai(text)

    # Identifiers stripped
    assert "32405891230018" not in sanitized
    assert "308904387" not in sanitized
    assert "+998 (90) 123-45-67" not in sanitized
    assert "info@supplier-uz.com" not in sanitized

    # Placeholders inserted
    assert "[PINFL]" in sanitized
    assert "[STIR]" in sanitized
    assert "[PHONE]" in sanitized
    assert "[EMAIL]" in sanitized

    # Commercial IT data preserved
    assert "Windows Server litsenziyasi" in sanitized


# -----------------------------------------------------------------------------
# 3. Admin Unmask Endpoint & Audit Log Integration
# -----------------------------------------------------------------------------
def test_admin_unmask_endpoint_and_audit_log(admin_user, sales_user, admin_cookie, sales_cookie):
    """Verify /api/admin/unmask/{stir} enforces RBAC and logs all unmask operations."""
    target_pinfl = "32405891230018"

    # 1. Non-admin (sales) gets 403 Forbidden
    client.cookies.set(SESSION_COOKIE_NAME, sales_cookie[SESSION_COOKIE_NAME])
    res_sales = client.get(f"/api/admin/unmask/{target_pinfl}")
    assert res_sales.status_code == 403
    assert "admin huquqi talab qilinadi" in res_sales.json()["detail"]

    # 2. Admin gets 200 OK
    client.cookies.set(SESSION_COOKIE_NAME, admin_cookie[SESSION_COOKIE_NAME])
    res_admin = client.get(f"/api/admin/unmask/{target_pinfl}")
    assert res_admin.status_code == 200
    data = res_admin.json()
    assert data["stir"] == target_pinfl
    assert data["unmasked"] is True

    # 3. Verify audit log entry was created
    with session_scope() as session:
        log_entry = (
            session.query(AuditLog)
            .filter(
                AuditLog.user_id == admin_user["id"],
                AuditLog.action == "unmask_pinfl",
                AuditLog.target_id == target_pinfl,
            )
            .first()
        )
        assert log_entry is not None
        assert log_entry.target_type == "stir"
        assert log_entry.details.get("original_length") == 14


# -----------------------------------------------------------------------------
# 4. XLSX Export Masking & Audit Logging
# -----------------------------------------------------------------------------
def test_export_xlsx_masks_jshshir_for_non_admin(
    admin_user, sales_user, admin_cookie, sales_cookie
):
    """Verify Excel export masks JSHSHIR for sales and audits unmasked admin export."""
    pinfl = f"5{int.from_bytes(os.urandom(6), 'big') % 10**13:013d}"
    masked_pinfl = f"*********{pinfl[-5:]}"
    raw_name = f"Test Individual Org {os.urandom(2).hex()}"

    with session_scope() as session:
        org = Organization(stir=pinfl, name_canonical=raw_name)
        session.add(org)
        session.flush()

        proc = Procedure(
            source="uzex",
            source_id=f"exp_{os.urandom(3).hex()}",
            title="Antivirus litsenziyasi 1 yillik",
            customer_org_id=org.id,
            completed_at=datetime.now(UTC) - timedelta(days=300),
            start_price=50_000_000.0,
        )
        session.add(proc)
        session.flush()

        cls = Classification(
            procedure_id=proc.id,
            is_it=True,
            confidence=0.95,
            category="SECURITY_SOFTWARE",
            brand="Kaspersky",
            method="rule",
        )
        session.add(cls)

        opp = RenewalOpportunity(
            procedure_id=proc.id,
            customer_org_id=org.id,
            category="SECURITY_SOFTWARE",
            brand="Kaspersky",
            lifecycle_months=12,
            last_purchase_at=datetime.now(UTC) - timedelta(days=300),
            expected_renewal_at=datetime.now(UTC) + timedelta(days=65),
            contact_by_at=datetime.now(UTC) + timedelta(days=5),
            score=85,
            computed_at=datetime.now(UTC),
        )
        session.add(opp)
        session.commit()
        proc_id = proc.id

    try:
        # 1. Sales export -> JSHSHIR masked in sheet
        client.cookies.set(SESSION_COOKIE_NAME, sales_cookie[SESSION_COOKIE_NAME])
        res_sales = client.get("/api/export/xlsx")
        assert res_sales.status_code == 200
        wb_sales = openpyxl.load_workbook(io.BytesIO(res_sales.content))
        values_sales = [
            str(c.value) for ws in wb_sales.worksheets for r in ws.iter_rows() for c in r if c.value
        ]

        # Full pinfl must NOT be in sales export; masked version MUST be present
        assert pinfl not in values_sales
        assert masked_pinfl in values_sales

        # 2. Admin export -> Full pinfl present, audit log recorded
        client.cookies.set(SESSION_COOKIE_NAME, admin_cookie[SESSION_COOKIE_NAME])
        res_admin = client.get("/api/export/xlsx")
        assert res_admin.status_code == 200
        wb_admin = openpyxl.load_workbook(io.BytesIO(res_admin.content))
        values_admin = [
            str(c.value) for ws in wb_admin.worksheets for r in ws.iter_rows() for c in r if c.value
        ]

        assert pinfl in values_admin

        # Verify audit log entry
        with session_scope() as session:
            entry = (
                session.query(AuditLog)
                .filter(
                    AuditLog.user_id == admin_user["id"],
                    AuditLog.action == "export_unmasked_pinfl",
                )
                .first()
            )
            assert entry is not None

    finally:
        with session_scope() as session:
            session.query(RenewalOpportunity).filter_by(procedure_id=proc_id).delete()
            session.query(Classification).filter_by(procedure_id=proc_id).delete()
            p = session.get(Procedure, proc_id)
            if p:
                session.delete(p)
            o = session.query(Organization).filter_by(stir=pinfl).first()
            if o:
                session.delete(o)
            session.commit()


# -----------------------------------------------------------------------------
# 5. Data Retention & Cleanup Tests
# -----------------------------------------------------------------------------
def test_parse_retention_period():
    """Verify retention period string parsing ('365d', '1y')."""
    assert parse_retention_period("365d") == 365
    assert parse_retention_period("30d") == 30
    assert parse_retention_period("1y") == 365
    assert parse_retention_period("2y") == 730

    with pytest.raises(ValueError, match="Invalid retention period"):
        parse_retention_period("invalid")

    with pytest.raises(ValueError, match="Invalid retention period"):
        parse_retention_period("365m")

    with pytest.raises(ValueError, match="must be positive"):
        parse_retention_period("0d")


def test_cleanup_old_snapshots(admin_user):
    """Verify cleanup_old_snapshots deletes snapshots beyond cutoff and records audit event."""
    with session_scope() as session:
        # 1 old snapshot (400 days old)
        old_snap = RawSnapshot(
            url="https://test.uz/old",
            fetched_at=datetime.now(UTC) - timedelta(days=400),
            sha256="old_sha_" + os.urandom(4).hex(),
            body=b"old_data",
        )
        # 1 fresh snapshot (10 days old)
        fresh_snap = RawSnapshot(
            url="https://test.uz/fresh",
            fetched_at=datetime.now(UTC) - timedelta(days=10),
            sha256="fresh_sha_" + os.urandom(4).hex(),
            body=b"fresh_data",
        )
        session.add_all([old_snap, fresh_snap])
        session.commit()
        old_id = old_snap.id
        fresh_id = fresh_snap.id

    try:
        with session_scope() as session:
            stats = cleanup_old_snapshots(session, older_than_days=365, user_id=admin_user["id"])
            session.commit()

        assert stats["deleted_snapshots"] >= 1
        assert stats["older_than_days"] == 365

        # Check DB: old snapshot must be gone, fresh snapshot must remain
        with session_scope() as session:
            assert session.get(RawSnapshot, old_id) is None
            assert session.get(RawSnapshot, fresh_id) is not None

            # Verify audit log entry
            audit = (
                session.query(AuditLog)
                .filter(
                    AuditLog.user_id == admin_user["id"],
                    AuditLog.action == "cleanup_snapshots",
                )
                .first()
            )
            assert audit is not None
            assert audit.details["older_than_days"] == 365
            assert audit.details["deleted"] >= 1

    finally:
        with session_scope() as session:
            f = session.get(RawSnapshot, fresh_id)
            if f:
                session.delete(f)
            session.commit()
